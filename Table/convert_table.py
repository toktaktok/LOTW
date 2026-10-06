"""
convert_table.py
────────────────────────────────────────────────────────────────
Excel (.xlsx / XML Spreadsheet 2003 .xml) → JSON 변환 스크립트

사용법:
  python convert_table.py           # Excel/ 내 모든 .xlsx, .xml 변환
  python convert_table.py Dialogue  # Dialogue.xlsx (또는 .xml) 만 변환

.xlsx 변환에는 openpyxl 이 필요합니다:  pip install openpyxl

출력 경로 (두 곳 동시):
  Table/Json/                          ← 소스 관리용
  Game/Assets/Project/Resources/Table/ ← Unity DataManager 로드 경로

표(AutoFilter) 규칙:
  - Excel에서 데이터 범위를 선택 후 [삽입 > 표]로 반드시 표를 정의해야 함 (B2 시작)
  - 표 1행 = 컬럼 헤더 (C# 필드명, 첫 글자만 자동 소문자 변환)
    예) DataId → dataId, SpeakerName → speakerName
  - 표 2행 = 자료형 (int, float, bool, string, int[]). Schema/{테이블}.json 과 같아야 함
  - 표 범위 바깥의 셀은 무시됨 (메모, 설명 등을 자유롭게 작성 가능)
  - int[] 컬럼은 쉼표로 구분 ("2,3" → [2, 3])
  - {스키마}_{분류}.xlsx 는 스키마를 같이 쓰는 분할 테이블 (예: Text_UI.xlsx → Schema/Text.json)

검증:
  모든 테이블을 읽어 Schema/ 규칙(자료형, 필수, 중복, 범위, 참조, '@키')으로 검사합니다.
  오류가 하나라도 있으면 목록을 출력하고 JSON 을 하나도 쓰지 않고 중단합니다.

C# RowData 생성:
  검증을 통과하면 스키마의 rowClass, columns 로 Scripts/Data/Table/Generated/{rowClass}.cs 를 씁니다.
  컬럼 = 필드 (camelCase, desc = 주석, default = 초기값). DataId 는 TableRowData.dataId.
  생성 클래스는 partial 이라 상수, 메서드는 Data/Table/{rowClass}.cs 에 둡니다.
  DataManager 가 읽을 테이블 목록도 Scripts/Core/Managers/Generated/DataManager.Tables.cs 로 씁니다
  (분할 테이블 Text_* 는 Localization 이 읽으므로 제외).
"""

import os
import sys
import re
import json
import html
import xml.etree.ElementTree as ET

from table_schema import Table, build_rows, load_schema, to_camel, validate

# ── 경로 설정 ──────────────────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_DIR  = os.path.join(SCRIPT_DIR, "Excel")
JSON_DIR   = os.path.join(SCRIPT_DIR, "Json")
SCHEMA_DIR = os.path.join(SCRIPT_DIR, "Schema")
UNITY_RES  = os.path.join(SCRIPT_DIR, "..", "Game", "Assets", "Project", "Resources", "Table")

OUTPUT_DIRS = [JSON_DIR, UNITY_RES]
ROWDATA_DIR = os.path.join(SCRIPT_DIR, "..", "Game", "Assets", "Project", "Scripts", "Data", "Table", "Generated")
REGISTRY_PATH = os.path.join(SCRIPT_DIR, "..", "Game", "Assets", "Project", "Scripts", "Core", "Managers", "Generated", "DataManager.Tables.cs")

# Excel XML Spreadsheet 2003 네임스페이스
NS  = "urn:schemas-microsoft-com:office:spreadsheet"
NSX = "urn:schemas-microsoft-com:office:excel"


# ── R1C1 범위 파싱 ─────────────────────────────────────────────
def parse_rc_range(raw: str):
    """
    R1C1:R6C6 형식을 (start_row, start_col, end_row, end_col) 로 변환.
    '=Dialogue!R1C1:R6C6' 처럼 시트 참조나 '=' 접두사가 있어도 처리합니다.
    """
    s = raw.strip().lstrip("=")
    if "!" in s:
        s = s.split("!")[-1]
    m = re.fullmatch(r"R(\d+)C(\d+):R(\d+)C(\d+)", s)
    if not m:
        return None
    return int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))


# ── 행 인덱스 구성 ─────────────────────────────────────────────
def build_row_map(table_elem) -> dict:
    """
    <Table> 안의 <Row> 요소를 행 번호(1-based) → element 로 매핑합니다.
    ss:Index 속성이 있는 행은 그 번호를 사용하고, 없으면 순서대로 증가합니다.
    """
    row_map = {}
    row_idx = 0
    for row_elem in table_elem.findall(f"{{{NS}}}Row"):
        idx_attr = row_elem.get(f"{{{NS}}}Index")
        row_idx  = int(idx_attr) if idx_attr else row_idx + 1
        row_map[row_idx] = row_elem
    return row_map


# ── 열 범위 내 셀 추출 ─────────────────────────────────────────
def get_cells_in_range(row_elem, start_col: int, end_col: int) -> dict:
    """
    행 요소에서 열 범위[start_col, end_col](1-based, 포함) 내 셀을
    {col_num: element} 딕셔너리로 반환합니다.
    ss:Index 속성으로 건너뛴 열도 정확하게 처리합니다.
    """
    cells = {}
    col_idx = 0
    for cell in row_elem.findall(f"{{{NS}}}Cell"):
        idx_attr = cell.get(f"{{{NS}}}Index")
        col_idx  = int(idx_attr) if idx_attr else col_idx + 1
        if start_col <= col_idx <= end_col:
            cells[col_idx] = cell
        if col_idx >= end_col:
            break
    return cells


# ── 셀 값 파싱 ────────────────────────────────────────────────
def parse_cell(cell) -> object:
    """셀 하나의 값을 Python 타입으로 반환합니다. 빈 셀이면 None."""
    data = cell.find(f"{{{NS}}}Data")
    if data is None or not data.text:
        return None
    text  = data.text.strip()
    dtype = data.get(f"{{{NS}}}Type", "String")
    if dtype == "Number":
        return float(text) if "." in text else int(text)
    if dtype == "Boolean":
        return text.lower() in ("1", "true", "yes")
    return text


# ── XML 파싱 메인 ─────────────────────────────────────────────
def parse_xml(filepath: str):
    """(그리드, 표 시작 행, 표 시작 열) 반환."""
    tree = ET.parse(filepath)
    root = tree.getroot()

    worksheet = root.find(f".//{{{NS}}}Worksheet")
    if worksheet is None:
        raise ValueError("Worksheet 요소를 찾을 수 없습니다.")

    table_elem = worksheet.find(f"{{{NS}}}Table")
    if table_elem is None:
        raise ValueError("Table 요소를 찾을 수 없습니다.")

    # ── 표(AutoFilter) 범위 읽기 ──────────────────────────────
    auto_filter = worksheet.find(f"{{{NSX}}}AutoFilter")
    if auto_filter is None:
        raise ValueError(
            "워크시트에 '표(Table)'가 정의되어 있지 않습니다.\n"
            "  Excel에서 데이터 범위를 선택한 뒤 [삽입 > 표]를 클릭하여\n"
            "  표를 만든 후 XML로 저장하세요."
        )

    raw_range = (
        auto_filter.get(f"{{{NSX}}}Range") or
        auto_filter.get("Range") or ""
    )
    if not raw_range:
        raise ValueError("AutoFilter에 Range 속성이 없습니다.")

    table_range = parse_rc_range(raw_range)
    if table_range is None:
        raise ValueError(f"표 범위 형식을 인식할 수 없습니다: '{raw_range}'")

    sr, sc, er, ec = table_range  # start_row, start_col, end_row, end_col (1-based)

    # ── 표 범위를 값 그리드로 변환 ────────────────────────────
    row_map = build_row_map(table_elem)
    grid = []
    for row_num in range(sr, er + 1):
        row_elem = row_map.get(row_num)
        cells = get_cells_in_range(row_elem, sc, ec) if row_elem is not None else {}
        grid.append([parse_cell(cells[c]) if c in cells else None for c in range(sc, ec + 1)])

    return grid, sr, sc


# ── XLSX 파싱 메인 ────────────────────────────────────────────
def parse_xlsx(filepath: str):
    """(그리드, 표 시작 행, 표 시작 열) 반환."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl 이 없습니다. 'pip install openpyxl' 후 다시 실행하세요.")

    sheet, table = find_xlsx_table(load_workbook(filepath, data_only=True))
    return read_xlsx_grid(sheet, table)


def find_xlsx_table(wb):
    """첫 번째 표가 있는 (시트, 표). 표가 없으면 ValueError."""
    for ws in wb.worksheets:
        if ws.tables:
            return ws, next(iter(ws.tables.values()))
    raise ValueError(
        "워크시트에 '표(Table)'가 정의되어 있지 않습니다.\n"
        "  Excel에서 데이터 범위를 선택한 뒤 [삽입 > 표]를 클릭하여 표를 만드세요."
    )


def read_xlsx_grid(sheet, table):
    """(그리드, 표 시작 행, 표 시작 열) 반환."""
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    grid = []
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row,
                               min_col=min_col, max_col=max_col, values_only=True):
        grid.append([normalize_xlsx_value(v) for v in row])

    return grid, min_row, min_col


def normalize_xlsx_value(value):
    """openpyxl 셀 값을 XML 파서와 같은 형태로 맞춥니다. 빈 셀이면 None."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


# ── 출력 ──────────────────────────────────────────────────────
def ensure_dirs():
    for d in OUTPUT_DIRS:
        os.makedirs(d, exist_ok=True)


def write_json(table_name: str, data: list):
    json_str = json.dumps(data, ensure_ascii=False, indent=2)
    for out_dir in OUTPUT_DIRS:
        path = os.path.join(out_dir, f"{table_name}.json")
        with open(path, "w", encoding="utf-8") as f:
            f.write(json_str)
        print(f"    → {os.path.normpath(path)}")


# ── C# RowData 생성 ───────────────────────────────────────────
def row_class_source(schema_name: str, schema: dict) -> str:
    """스키마 컬럼으로 partial RowData 클래스 소스를 만듭니다. 스키마 type 은 C# 자료형 이름과 같습니다."""
    lines = [
        f"// 자동 생성 파일. Table/Schema/{schema_name}.json 을 고치고 ConvertTable.bat 을 실행하세요.",
        "using System;",
        "",
        "namespace Project.Scripts.Data.Table",
        "{",
        "    [Serializable]",
        f"    public partial class {schema['rowClass']} : TableRowData",
        "    {",
    ]
    for col in schema["columns"]:
        if col["name"] == "DataId":
            continue  # TableRowData.dataId
        if col.get("desc"):
            lines.append(f"        /// <summary>{html.escape(' '.join(col['desc'].split()), quote=False)}</summary>")
        default = ""
        if "default" in col:
            default = f" = {json.dumps(col['default'], ensure_ascii=False)}{'f' if col['type'] == 'float' else ''}"
        lines.append(f"        public {col['type']} {to_camel(col['name'])}{default};")
    lines += ["    }", "}", ""]
    return "\n".join(lines)


def registry_source(tables: dict) -> str:
    """DataManager 가 dataId 로 캐싱할 테이블: rowClass 가 있고 분할되지 않은 테이블 (Text_* 제외)."""
    rows = sorted((t.name, t.schema["rowClass"]) for t in tables.values()
                  if t.schema.get("rowClass") and t.name == t.schema_name)
    lines = [
        "// 자동 생성 파일. Table/Excel 과 Table/Schema 로 ConvertTable.bat 이 만듭니다.",
        "namespace Project.Scripts.Core.Managers",
        "{",
        "    public partial class DataManager",
        "    {",
        "        private void LoadSchemaTables()",
        "        {",
    ]
    lines += [f'            LoadTable<Project.Scripts.Data.Table.{row_class}>("{name}");' for name, row_class in rows]
    lines += ["        }", "    }", "}", ""]
    return "\n".join(lines)


def write_generated(path: str, source: str):
    """내용이 같으면 건드리지 않습니다 (Unity 재컴파일 방지)."""
    if os.path.isfile(path):
        with open(path, encoding="utf-8") as f:
            if f.read() == source:
                return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(source)
    print(f"    → {os.path.normpath(path)}")


def write_row_classes(tables: dict):
    """rowClass 가 있는 스키마마다 RowData 파일을, 그리고 DataManager 테이블 목록을 씁니다."""
    schemas = {t.schema_name: t.schema for t in tables.values() if t.schema.get("rowClass")}
    for schema_name, schema in sorted(schemas.items()):
        write_generated(os.path.join(ROWDATA_DIR, f"{schema['rowClass']}.cs"), row_class_source(schema_name, schema))
    write_generated(REGISTRY_PATH, registry_source(tables))


# ── 읽기 + 검증 ───────────────────────────────────────────────
def read_tables(grids: dict = None):
    """
    Excel/ 의 모든 테이블을 읽고 검증해 (tables, errors) 를 반환합니다.
    grids 에 {테이블 이름: (그리드, 시작 행, 시작 열)} 이 있으면 파일 대신 그 값을 씁니다 (table_edit.py).
    """
    parsers = {".xlsx": parse_xlsx, ".xml": parse_xml}
    grids = grids or {}

    # '~$' 로 시작하는 파일은 Excel 이 열려 있을 때 생기는 잠금 파일
    table_files = [
        f for f in os.listdir(EXCEL_DIR)
        if os.path.splitext(f)[1].lower() in parsers
        and not f.startswith("~$")
    ]

    # 참조 검사를 위해 필터와 관계없이 모든 테이블을 읽음
    tables, errors = {}, []
    for table_file in table_files:
        table_name, ext = os.path.splitext(table_file)
        table = Table(table_name, table_file)
        table.schema = load_schema(SCHEMA_DIR, table.schema_name)
        if table.schema is None:
            errors.append(f"{table_file}: Schema/{table.schema_name}.json 이 없습니다")
            continue
        try:
            source = grids.get(table_name) or parsers[ext.lower()](os.path.join(EXCEL_DIR, table_file))
            grid, table.origin_row, table.origin_col = source
        except Exception as e:
            errors.append(f"{table_file}: {e}")
            continue
        build_rows(table, grid, errors)
        tables[table_name] = table

    if not errors:
        errors = validate(tables)
    return tables, errors


# ── 메인 ──────────────────────────────────────────────────────
def convert(filter_name: str = "") -> bool:
    ensure_dirs()

    if not os.path.isdir(EXCEL_DIR):
        print(f"[오류] {EXCEL_DIR} 디렉토리가 없습니다.")
        return False

    tables, errors = read_tables()
    if not tables and not errors:
        print("[경고] 변환할 테이블 파일이 없습니다.")
        return True

    if errors:
        print(f"[검증 실패] {len(errors)}개 오류. JSON 을 쓰지 않고 중단합니다.\n")
        for e in errors:
            print(f"  ✗ {e}")
        return False

    write_row_classes(tables)
    targets = [t for t in tables.values() if not filter_name or filter_name.lower() in t.file.lower()]
    for table in targets:
        print(f"  [{table.file}] {len(table.rows)}개 행")
        write_json(table.name, [row for _, row in table.rows])
    print(f"\n결과: {len(targets)}개 테이블 변환 완료")
    return True


if __name__ == "__main__":
    target  = sys.argv[1] if len(sys.argv) > 1 else ""
    success = convert(target)
    sys.exit(0 if success else 1)
