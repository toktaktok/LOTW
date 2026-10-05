"""
table_edit.py
────────────────────────────────────────────────────────────────
Unity Dialogue 에디터가 xlsx 표의 행을 추가, 수정, 삭제할 때 부르는 스크립트입니다.

사용법:
  python table_edit.py edits.json

edits.json (필드명은 JSON 과 같은 camelCase):
  [{"table": "Dialogue", "key": "dataId",
    "upsert": [{"dataId": 1, "text": "@dialogue.1", "nextId": -1, "choiceIds": [2, 3]}],
    "delete": [4]},
   {"table": "Text_Dialogue", "key": "key",
    "upsert": [{"key": "dialogue.1", "ko": "안녕"}], "delete": ["dialogue.4"]}]

순서: 잠금 파일 확인 → 메모리에서 편집 → 전체 테이블 검증 → 통과하면 xlsx 저장 → JSON 변환.
검증에 실패하면 아무 파일도 쓰지 않습니다.

쓰기 규칙:
  - upsert 에 없는 필드는 그대로 둡니다 (예: Text 의 En).
  - "" 와 [] 는 빈 셀, int[] 는 "2,3" 문자열로 씁니다.
  - key 가 dataId 인 표는 새 행을 DataId 순서 자리에 넣고, 아니면 끝에 붙입니다.
    dataId 없이 들어온 새 행은 마지막 DataId + 1 을 받습니다.
  - 표 범위를 행 수에 맞게 늘리거나 줄이고, string / int[] 열은 텍스트 서식으로 둡니다.
"""

import json
import os
import sys

import convert_table
from table_schema import to_camel

TEXT_TYPES = ("string", "int[]")


def to_cell(value):
    if value is None or value == "" or value == []:
        return None
    if isinstance(value, list):
        return ",".join(str(v) for v in value)
    return value


def apply_edit(sheet, table, edit):
    """표 하나에 upsert / delete 를 반영합니다 (메모리에서만)."""
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    cols = range(min_col, max_col + 1)
    fields = [to_camel(str(sheet.cell(min_row, c).value or "").strip()) for c in cols]
    types = [str(sheet.cell(min_row + 1, c).value or "").strip() for c in cols]
    key = fields.index(edit["key"])
    id_col = fields.index("dataId") if "dataId" in fields else None

    deleted = set(edit.get("delete", []))
    rows = [[sheet.cell(r, c).value for c in cols] for r in range(min_row + 2, max_row + 1)]
    rows = [r for r in rows if r[key] not in deleted]

    for new in edit.get("upsert", []):
        row = next((r for r in rows if r[key] == new[edit["key"]]), None)
        if row is None:
            row = [None] * len(fields)
            if id_col is not None and new.get("dataId") is None:
                row[id_col] = max([r[id_col] for r in rows if isinstance(r[id_col], int)], default=0) + 1
            if edit["key"] == "dataId":
                at = next((i for i, r in enumerate(rows) if isinstance(r[key], int) and r[key] > new["dataId"]), len(rows))
                rows.insert(at, row)
            else:
                rows.append(row)
        for field, value in new.items():
            if field in fields:
                row[fields.index(field)] = to_cell(value)

    for offset, values in enumerate(rows):
        for i, c in enumerate(cols):
            # sheet.cell(r, c, None) 은 값을 지우지 않으므로 직접 대입
            cell = sheet.cell(min_row + 2 + offset, c)
            cell.value = values[i]
            cell.number_format = "@" if types[i] in TEXT_TYPES else "General"
    end_row = min_row + 1 + len(rows)
    for r in range(end_row + 1, max_row + 1):
        for c in cols:
            sheet.cell(r, c).value = None

    table.ref = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(end_row, max_col).coordinate}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref


def edit_tables(edits) -> bool:
    from openpyxl import load_workbook

    locked = [e["table"] for e in edits if os.path.exists(os.path.join(convert_table.EXCEL_DIR, f"~${e['table']}.xlsx"))]
    if locked:
        print(f"[저장 실패] 엑셀에서 열려 있습니다: {', '.join(locked)}. 엑셀을 닫고 다시 저장하세요.")
        return False

    books, grids = {}, {}
    for edit in edits:
        path = os.path.join(convert_table.EXCEL_DIR, f"{edit['table']}.xlsx")
        wb = load_workbook(path)
        sheet, table = convert_table.find_xlsx_table(wb)
        apply_edit(sheet, table, edit)
        books[path] = wb
        grids[edit["table"]] = convert_table.read_xlsx_grid(sheet, table)

    _, errors = convert_table.read_tables(grids)
    if errors:
        print(f"[검증 실패] {len(errors)}개 오류. 엑셀을 저장하지 않았습니다.\n")
        for e in errors:
            print(f"  ✗ {e}")
        return False

    for path, wb in books.items():
        wb.save(path)
        print(f"  [저장] {os.path.basename(path)}")
    return convert_table.convert()


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        success = edit_tables(json.load(f))
    sys.exit(0 if success else 1)
