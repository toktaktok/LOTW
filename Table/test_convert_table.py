"""
convert_table / table_schema 검증 테스트.
실행: python -m unittest test_convert_table   (Table 폴더에서)
"""

import json
import os
import tempfile
import unittest

import convert_table
from table_schema import Table, build_rows, validate

SCHEMAS = {
    "Dialogue": {"columns": [
        {"name": "DataId", "type": "int", "required": True, "unique": True, "min": 1},
        {"name": "Text", "type": "string", "textKey": True},
        {"name": "NextId", "type": "int", "ref": "Dialogue.DataId", "allow": [-1]},
        {"name": "ChoiceIds", "type": "int[]", "ref": "Dialogue.DataId", "maxCount": 2},
    ]},
    "Text": {"columns": [
        {"name": "DataId", "type": "int", "required": True, "unique": True},
        {"name": "Key", "type": "string", "required": True, "unique": True},
    ]},
}
HEADER = ["DataId", "Text", "NextId", "ChoiceIds"]
TYPES = ["int", "string", "int", "int[]"]
TEXT_GRID = [["DataId", "Key"], ["int", "string"], [1, "hello"]]


def run(dialogue_grid):
    tables, errors = {}, []
    for name, grid in (("Dialogue", dialogue_grid), ("Text", TEXT_GRID)):
        table = Table(name, f"{name}.xlsx")
        table.origin_row, table.origin_col = 2, 2
        table.schema = SCHEMAS[name]
        build_rows(table, grid, errors)
        tables[name] = table
    return tables, errors + validate(tables)


class ValidateTests(unittest.TestCase):
    def test_valid_rows_convert(self):
        tables, errors = run([HEADER, TYPES, [1, "@hello", 2, None], [2, "@hello", -1, "1,2"], [None] * 4])
        self.assertEqual(errors, [])
        self.assertEqual([r for _, r in tables["Dialogue"].rows],
                         [{"dataId": 1, "text": "@hello", "nextId": 2}, {"dataId": 2, "text": "@hello", "nextId": -1, "choiceIds": [1, 2]}])

    def test_each_rule_reports_error(self):
        cases = {
            "자료형 'string'": [HEADER, ["string", "string", "int", "int[]"], [1, "a", -1, None]],
            "스키마에 없는 컬럼": [HEADER + ["Extra"], TYPES + ["int"], [1, "a", -1, None, 0]],
            "읽을 수 없습니다": [HEADER, TYPES, [1, "a", "x", None]],
            "값이 필요합니다": [HEADER, TYPES, [None, "a", -1, None]],
            "중복": [HEADER, TYPES, [1, "a", -1, None], [1, "b", -1, None]],
            "최소값": [HEADER, TYPES, [0, "a", -1, None]],
            "Dialogue.DataId 에 없습니다": [HEADER, TYPES, [1, "a", 9, None]],
            "최대 2개": [HEADER, TYPES, [1, "a", -1, "1,1,1"]],
            "Text 키 'missing' 가 없습니다": [HEADER, TYPES, [1, "@missing", -1, None]],
            "Text 키('@키')여야": [HEADER, TYPES, [1, "literal", -1, None]],
        }
        for expected, grid in cases.items():
            _, errors = run(grid)
            self.assertTrue(any(expected in e for e in errors), f"{expected}: {errors}")

    def test_error_points_at_cell(self):
        _, errors = run([HEADER, TYPES, [1, "a", "x", None]])
        self.assertIn("Dialogue.xlsx D4", errors[0])
        _, errors = run([HEADER, TYPES, [1, "@hello", 9, None]])
        self.assertIn("Dialogue.xlsx D4(NextId)", errors[0])


class CommandRefTests(unittest.TestCase):
    """Dialogue Conditions/Actions 의 아이템 ID 가 Item 테이블에 있는지 (commandRefs)."""

    def run_commands(self, actions):
        schemas = {
            "Dialogue": {"columns": [
                {"name": "DataId", "type": "int"},
                {"name": "Actions", "type": "string",
                 "commandRefs": {"giveItem": "Item.ItemId", "item": "Item.ItemId", "startQuest": "Quest.DataId"}},
            ]},
            "Item": {"columns": [{"name": "ItemId", "type": "string"}]},
            "Quest": {"columns": [{"name": "DataId", "type": "int"}]},
        }
        grids = {
            "Dialogue": [["DataId", "Actions"], ["int", "string"], [1, actions]],
            "Item": [["ItemId"], ["string"], ["rose"]],
            "Quest": [["DataId"], ["int"], [101]],
        }
        tables, errors = {}, []
        for name, grid in grids.items():
            table = Table(name, f"{name}.xlsx")
            table.origin_row, table.origin_col = 2, 2
            table.schema = schemas[name]
            build_rows(table, grid, errors)
            tables[name] = table
        return errors + validate(tables)

    def test_known_item_passes(self):
        self.assertEqual(self.run_commands("giveItem:rose=2;setFlag:tulip;!item:rose>=1"), [])

    def test_missing_item_reports_error(self):
        errors = self.run_commands("setFlag:got;giveitem:tulip=1")
        self.assertEqual(len(errors), 1)
        self.assertIn("'tulip' 가 Item.ItemId 에 없습니다", errors[0])

    def test_int_id_command(self):
        self.assertEqual(self.run_commands("startQuest:101"), [])
        errors = self.run_commands("startQuest:102")
        self.assertEqual(len(errors), 1)
        self.assertIn("'102' 가 Quest.DataId 에 없습니다", errors[0])


class RefWhenTests(unittest.TestCase):
    """Sequence Param 처럼 다른 컬럼 값에 따라 참조를 검사 (refWhen)."""

    def run_steps(self, rows):
        schemas = {
            "Sequence": {"columns": [
                {"name": "Type", "type": "string"},
                {"name": "Param", "type": "string", "refWhen": {"column": "Type", "value": "dialogue", "ref": "Dialogue.DataId"}},
            ]},
            "Dialogue": {"columns": [{"name": "DataId", "type": "int"}]},
        }
        grids = {
            "Sequence": [["Type", "Param"], ["string", "string"]] + rows,
            "Dialogue": [["DataId"], ["int"], [9000]],
        }
        tables, errors = {}, []
        for name, grid in grids.items():
            table = Table(name, f"{name}.xlsx")
            table.schema = schemas[name]
            build_rows(table, grid, errors)
            tables[name] = table
        return errors + validate(tables)

    def test_checks_only_matching_type(self):
        self.assertEqual(self.run_steps([["Dialogue", "9000"], ["scene", "Plaza"]]), [])
        errors = self.run_steps([["dialogue", "9001"]])
        self.assertEqual(len(errors), 1)
        self.assertIn("9001 이 Dialogue.DataId 에 없습니다", errors[0])

    def test_empty_param_on_matching_type(self):
        self.assertEqual(self.run_steps([["scene", None]]), [])
        errors = self.run_steps([["dialogue", None]])
        self.assertEqual(len(errors), 1)
        self.assertIn("값이 필요합니다", errors[0])


class SplitTableTests(unittest.TestCase):
    """Text_UI, Text_Item 처럼 Schema/Text.json 을 같이 쓰는 분할 테이블."""

    SCHEMA = {"columns": [
        {"name": "DataId", "type": "int", "required": True, "unique": True},
        {"name": "Key", "type": "string", "unique": "group", "pattern": "[a-z0-9_]+(\\.[a-z0-9_]+)+", "prefixByPart": True},
    ]}

    def run_split(self, ui_rows, item_rows, dialogue_text="@ui.talk"):
        tables, errors = {}, []
        grids = {
            "Text_UI": [["DataId", "Key"], ["int", "string"]] + ui_rows,
            "Text_Item": [["DataId", "Key"], ["int", "string"]] + item_rows,
            "Dialogue": [HEADER, TYPES, [1, dialogue_text, -1, None]],
        }
        for name, grid in grids.items():
            table = Table(name, f"{name}.xlsx")
            table.schema = SCHEMAS["Dialogue"] if name == "Dialogue" else self.SCHEMA
            build_rows(table, grid, errors)
            tables[name] = table
        return errors + validate(tables)

    def test_data_id_is_per_file_and_text_key_spans_group(self):
        errors = self.run_split([[1, "ui.talk"]], [[1, "item.rose.name"]], "@item.rose.name")
        self.assertEqual(errors, [])

    def test_group_rules(self):
        cases = {
            "Text_UI.xlsx 3행과 중복": ([[1, "ui.talk"]], [[1, "ui.talk"]]),
            "'item.' 로 시작해야": ([[1, "ui.talk"]], [[1, "ui.rose"]]),
            "형식": ([[1, "ui.Talk"]], []),
        }
        for expected, (ui_rows, item_rows) in cases.items():
            errors = self.run_split(ui_rows, item_rows)
            self.assertTrue(any(expected in e for e in errors), f"{expected}: {errors}")


class RowClassTests(unittest.TestCase):
    """스키마 → C# RowData 생성."""

    def test_source_has_fields_defaults_and_escaped_desc(self):
        schema = {"rowClass": "DialogueData", "columns": [
            {"name": "DataId", "type": "int"},
            {"name": "NextId", "type": "int", "default": -1, "desc": "a<b"},
            {"name": "Speed", "type": "float", "default": 1.5},
            {"name": "ChoiceIds", "type": "int[]"},
        ]}
        source = convert_table.row_class_source("Dialogue", schema)
        self.assertIn("public partial class DialogueData : TableRowData", source)
        self.assertNotIn("dataId", source)
        self.assertIn("/// <summary>a&lt;b</summary>\n        public int nextId = -1;", source)
        self.assertIn("public float speed = 1.5f;", source)
        self.assertIn("public int[] choiceIds;", source)

    def test_generated_files_match_schemas(self):
        """스키마를 고치고 변환을 안 돌리면 실패합니다."""
        for file in os.listdir(convert_table.SCHEMA_DIR):
            name = os.path.splitext(file)[0]
            with open(os.path.join(convert_table.SCHEMA_DIR, file), encoding="utf-8") as f:
                schema = json.load(f)
            if "rowClass" not in schema:
                continue
            with open(os.path.join(convert_table.ROWDATA_DIR, f"{schema['rowClass']}.cs"), encoding="utf-8") as f:
                self.assertEqual(f.read(), convert_table.row_class_source(name, schema), file)


class RepoUpToDateTests(unittest.TestCase):
    """Excel 을 고치고 ConvertTable.bat 을 안 돌리면 실패합니다 (커밋된 JSON, DataManager 목록 비교)."""

    @classmethod
    def setUpClass(cls):
        cls.tables, cls.errors = convert_table.read_tables()

    def test_excel_is_valid(self):
        self.assertEqual(self.errors, [])

    def test_json_matches_excel(self):
        for table in self.tables.values():
            expected = [row for _, row in table.rows]
            for out_dir in convert_table.OUTPUT_DIRS:
                with open(os.path.join(out_dir, f"{table.name}.json"), encoding="utf-8") as f:
                    self.assertEqual(json.load(f), expected, f"{out_dir} {table.name}.json")

    def test_registry_matches_tables(self):
        with open(convert_table.REGISTRY_PATH, encoding="utf-8") as f:
            self.assertEqual(f.read(), convert_table.registry_source(self.tables))


class ConvertTests(unittest.TestCase):
    def test_error_stops_without_writing_json(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table as XlTable

        with tempfile.TemporaryDirectory() as tmp:
            excel, schema, out = (os.path.join(tmp, d) for d in ("Excel", "Schema", "Out"))
            os.makedirs(excel)
            os.makedirs(schema)
            for name, grid in (("Dialogue", [HEADER, TYPES, [1, "a", 9, None]]), ("Text", TEXT_GRID)):
                with open(os.path.join(schema, f"{name}.json"), "w", encoding="utf-8") as f:
                    json.dump(SCHEMAS[name], f)
                wb = Workbook()
                for r, values in enumerate(grid, start=2):
                    for c, v in enumerate(values, start=2):
                        wb.active.cell(r, c, v)
                ref = f"B2:{chr(65 + len(grid[0]))}{len(grid) + 1}"
                wb.active.add_table(XlTable(displayName=name, ref=ref))
                wb.save(os.path.join(excel, f"{name}.xlsx"))

            saved = convert_table.EXCEL_DIR, convert_table.SCHEMA_DIR, convert_table.OUTPUT_DIRS
            convert_table.EXCEL_DIR, convert_table.SCHEMA_DIR, convert_table.OUTPUT_DIRS = excel, schema, [out]
            try:
                self.assertFalse(convert_table.convert())
            finally:
                convert_table.EXCEL_DIR, convert_table.SCHEMA_DIR, convert_table.OUTPUT_DIRS = saved
            self.assertEqual(os.listdir(out), [])


if __name__ == "__main__":
    unittest.main()
