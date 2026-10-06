"""
table_edit 테스트.
실행: python -m unittest test_table_edit   (Table 폴더에서)
"""

import json
import os
import tempfile
import unittest

import convert_table
import table_edit
from test_convert_table import HEADER, SCHEMAS, TYPES

DIALOGUE = [HEADER, TYPES, [1, "@dialogue.1", 2, None], [2, "@dialogue.2", -1, None], [5, "@dialogue.1", -1, "1"]]
TEXT = [["DataId", "Key", "Ko"], ["int", "string", "string"], [1, "dialogue.1", "hi"], [2, "dialogue.2", "bye"]]
TEXT_SCHEMA = {"columns": [
    {"name": "DataId", "type": "int", "required": True, "unique": True},
    {"name": "Key", "type": "string", "required": True, "unique": "group"},
    {"name": "Ko", "type": "string", "required": True},
]}


class TableEditTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table as XlTable

        self.tmp = tempfile.TemporaryDirectory()
        excel, schema, out = (os.path.join(self.tmp.name, d) for d in ("Excel", "Schema", "Out"))
        for d in (excel, schema):
            os.makedirs(d)
        for name, grid, sch in (("Dialogue", DIALOGUE, SCHEMAS["Dialogue"]), ("Text_Dialogue", TEXT, TEXT_SCHEMA)):
            with open(os.path.join(schema, f"{name.split('_')[0]}.json"), "w", encoding="utf-8") as f:
                json.dump(sch, f)
            wb = Workbook()
            for r, values in enumerate(grid, start=2):
                for c, v in enumerate(values, start=2):
                    wb.active.cell(r, c, v)
            wb.active["B20"] = "note"
            wb.active.add_table(XlTable(displayName=name, ref=f"B2:{chr(65 + len(grid[0]))}{len(grid) + 1}"))
            wb.save(os.path.join(excel, f"{name}.xlsx"))

        self.saved = convert_table.EXCEL_DIR, convert_table.SCHEMA_DIR, convert_table.OUTPUT_DIRS, convert_table.REGISTRY_PATH
        convert_table.EXCEL_DIR, convert_table.SCHEMA_DIR, convert_table.OUTPUT_DIRS = excel, schema, [out]
        convert_table.REGISTRY_PATH = os.path.join(out, "DataManager.Tables.cs")
        self.excel, self.out = excel, out

    def tearDown(self):
        convert_table.EXCEL_DIR, convert_table.SCHEMA_DIR, convert_table.OUTPUT_DIRS, convert_table.REGISTRY_PATH = self.saved
        self.tmp.cleanup()

    def read(self, name):
        from openpyxl import load_workbook
        sheet, table = convert_table.find_xlsx_table(load_workbook(os.path.join(self.excel, f"{name}.xlsx")))
        return convert_table.read_xlsx_grid(sheet, table)[0], table.ref, sheet

    def test_upsert_delete_and_convert(self):
        ok = table_edit.edit_tables([
            {"table": "Dialogue", "key": "dataId",
             "upsert": [{"dataId": 3, "text": "@dialogue.3", "nextId": -1, "choiceIds": []},
                        {"dataId": 1, "text": "@dialogue.1", "nextId": 3, "choiceIds": [2, 3]}],
             "delete": [5]},
            {"table": "Text_Dialogue", "key": "key",
             "upsert": [{"key": "dialogue.3", "ko": "new"}, {"key": "dialogue.2", "ko": "bye!"}]},
        ])
        self.assertTrue(ok)

        grid, ref, sheet = self.read("Dialogue")
        self.assertEqual(grid[2:], [[1, "@dialogue.1", 3, "2,3"], [2, "@dialogue.2", -1, None], [3, "@dialogue.3", -1, None]])
        self.assertEqual(ref, "B2:E6")
        self.assertEqual(sheet["E4"].number_format, "@")
        self.assertEqual(sheet["B20"].value, "note")

        grid, ref, _ = self.read("Text_Dialogue")
        self.assertEqual(grid[2:], [[1, "dialogue.1", "hi"], [2, "dialogue.2", "bye!"], [3, "dialogue.3", "new"]])

        with open(os.path.join(self.out, "Dialogue.json"), encoding="utf-8") as f:
            self.assertEqual([r["dataId"] for r in json.load(f)], [1, 2, 3])

    def test_validation_error_writes_nothing(self):
        path = os.path.join(self.excel, "Dialogue.xlsx")
        before = open(path, "rb").read()
        ok = table_edit.edit_tables([{"table": "Dialogue", "key": "dataId", "upsert": [{"dataId": 2, "nextId": 99}]}])
        self.assertFalse(ok)
        self.assertEqual(open(path, "rb").read(), before)
        self.assertFalse(os.path.exists(self.out))

    def test_lock_file_blocks_save(self):
        open(os.path.join(self.excel, "~$Dialogue.xlsx"), "w").close()
        ok = table_edit.edit_tables([{"table": "Dialogue", "key": "dataId", "delete": [5]}])
        self.assertFalse(ok)
        self.assertEqual(len(self.read("Dialogue")[0]), 5)


if __name__ == "__main__":
    unittest.main()
